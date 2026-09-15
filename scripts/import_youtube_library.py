"""Import only the Youtube sheet from the editorial workbook into the web app."""

from datetime import date, datetime, time
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('/Users/gasparclemente/Downloads/Temas mensagens Apóstolo Marcelino.xlsx')
OUTPUT = ROOT / 'data' / 'youtube-teachings.json'

def value_as_text(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, time):
        return value.strftime('%H:%M:%S')
    return value.strip() if isinstance(value, str) else value

workbook = load_workbook(SOURCE, read_only=True, data_only=True)
worksheet = workbook['Youtube']
records = []

for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
    published_at, title, biblical_reference, service, platform, url, starts_at, ends_at, description, *_ = row
    if not title or not url:
        continue
    records.append({
        'id': f'youtube-{row_number}',
        'publishedAt': value_as_text(published_at),
        'title': value_as_text(title),
        'biblicalReference': value_as_text(biblical_reference),
        'service': value_as_text(service),
        'platform': value_as_text(platform),
        'url': value_as_text(url),
        'startsAt': value_as_text(starts_at),
        'endsAt': value_as_text(ends_at),
        'description': value_as_text(description),
        'sourceSheet': 'Youtube'
    })

records.sort(key=lambda item: item['publishedAt'] or '', reverse=True)
OUTPUT.write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(f'Imported {len(records)} records from the Youtube sheet into {OUTPUT.relative_to(ROOT)}')
