"""Turn the team's general list of services into the church directory.

The workbook has one row per service: Kifica appears three times, once for
Thursday, once for Saturday and once for Sunday. The app wants one record per
place with all its services, so this script groups the rows, keeps every
service, and writes:

    data/igrejas.json                   what the app reads when the database
                                        cannot be reached, and what the
                                        migration below is built from
    data/importacao-igrejas.md          every decision taken on the way, for
                                        the team to check
    supabase/migrations/<name>.sql      optional: brings an existing database
                                        to the same state (--migracao)

    python3 scripts/import_churches.py <workbook.xlsx> [--migracao supabase/migrations/<name>.sql]

Nothing is invented. Addresses are the workbook's own text; a place whose
"address" is only its own name is left without one rather than given one.
Where the rows of one place disagree (a phone number, the spelling of a name)
the majority wins and the disagreement is written down. The few corrections
made are listed in REGION_FIXES and SAME_PLACE, and reported.

The national seats are the one thing the workbook does not say. SEATS holds the
proposal sent to the team; the panel changes it in a second.
"""

import argparse
import json
import re
import unicodedata
from collections import Counter, OrderedDict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
SHEET = 'Todos os Cultos'
SOURCE = 'Lista geral de cultos da equipa (setembro de 2026)'

WEEKDAYS = {'domingo': 0, 'segunda-feira': 1, 'terca-feira': 2, 'quarta-feira': 3,
            'quinta-feira': 4, 'sexta-feira': 5, 'sabado': 6}

COUNTRY_CODES = {
    'Angola': 'AO', 'Brasil': 'BR', 'Portugal': 'PT', 'Moçambique': 'MZ', 'Alemanha': 'DE',
    'Inglaterra': 'GB', 'França': 'FR', 'São Tomé e Príncipe': 'ST',
    'Estados Unidos da América': 'US', 'Canadá': 'CA', 'África do Sul': 'ZA', 'Suíça': 'CH',
    'Noruega': 'NO', 'Congo': 'CD', 'Tanzânia': 'TZ', 'Quénia': 'KE', 'Senegal': 'SN', 'Polónia': 'PL',
}

# The ministry's own abbreviations (src/roles.js). The workbook writes the same
# rank several ways; the app shows one.
RANKS = {
    'bp.': 'Bp.', 'bp': 'Bp.', 'bispo': 'Bp.',
    'bp. aux.': 'Bp. Aux.', 'bispo auxiliar': 'Bp. Aux.',
    'pr.': 'Pr.', 'pr': 'Pr.', 'pastor': 'Pr.',
    'pr. aux.': 'Pr. Aux.', 'pr.aux.': 'Pr. Aux.', 'pr. auxiliar': 'Pr. Aux.', 'pastor auxiliar': 'Pr. Aux.',
}
RANK_ORDER = ['Ap.', 'Bp.', 'Bp. Aux.', 'Pr.', 'Pr. Aux.']

# Spelling of provinces and states: typing slips and one abbreviation.
REGION_FIXES = {
    'Icole e Bengo': 'Icolo e Bengo',
    'Góias': 'Goiás',
    'MG': 'Minas Gerais',
    'Maputo cidade': 'Maputo Cidade',
}

# Rows the workbook lists under two names that are one place: the same
# responsible, the same number, the name of one containing the other.
SAME_PLACE = {
    ('Angola', 'Cazenga'): 'Cazenga - Kalawenda',
    ('Angola', 'Quibala'): 'Quibala - Cambango',
    ('Brasil', 'São Paulo'): 'São Paulo - Sede Estadual',
    ('Brasil', 'Sede Estadual'): 'São Paulo - Sede Estadual',
    ('Brasil', 'Baia'): 'Terra Nova',
}

# The proposal for the seats, to be confirmed by the team. The world seat is
# Kifica: first in the list, the only address written in full, and the temple
# the ministry itself calls its first. For the countries, the rule is the place
# of the highest-ranking responsible, the first listed on a tie.
SEATS = {
    'ao-kifica': 'mundial',
    'br-alto-garcas': 'nacional',
    'pt-pontinha': 'nacional',
    'mz-lualane': 'nacional',
    'de-wolfsburg': 'nacional',
    'gb-leeds': 'nacional',
    'fr-paris': 'nacional',
    'st-neves': 'nacional',
}


def clean(value):
    if value is None:
        return None
    text = re.sub(r'\s+', ' ', str(value)).strip()
    return text or None


def fold(value):
    return unicodedata.normalize('NFD', value or '').encode('ascii', 'ignore').decode().lower().strip()


def slug(value):
    return re.sub(r'[^a-z0-9]+', '-', fold(value)).strip('-')


def digits(phone):
    return re.sub(r'\D', '', phone or '')


def rank_of(cargo):
    rank = RANKS.get(fold(cargo).replace('  ', ' '))
    if not rank:
        raise SystemExit(f'Cargo desconhecido: {cargo!r}. Acrescente-o a RANKS.')
    return rank


def majority(values, prefer=None):
    """The most frequent value; on a tie, `prefer` if it is among the tied."""
    counts = Counter(value for value in values if value)
    if not counts:
        return None
    top = max(counts.values())
    tied = [value for value, count in counts.items() if count == top]
    if prefer in tied:
        return prefer
    return tied[0]


def read_rows(workbook):
    sheet = load_workbook(workbook, read_only=True, data_only=True)[SHEET]
    rows = sheet.iter_rows(values_only=True)
    header = [clean(cell) for cell in next(rows)]
    return [dict(zip(header, row)) for row in rows if any(cell not in (None, '') for cell in row)]


def build(workbook):
    notes = {'merged': [], 'regions': [], 'names': [], 'phones': [], 'leaders': [], 'addresses': []}
    places = OrderedDict()

    for number, row in enumerate(read_rows(workbook), start=2):
        country = clean(row['País'])
        online = fold(row['Modalidade']) == 'online'
        written = clean(row['Localidade/Igreja'])
        locality = written
        if not online and (country, locality) in SAME_PLACE:
            target = SAME_PLACE[(country, locality)]
            notes['merged'].append(f'Linha {number}: «{locality}» ({country}) junta-se a «{target}».')
            locality = target
        key = (country, 'online') if online else (country, fold(locality))
        place = places.setdefault(key, {'country': country, 'online': online, 'rows': []})
        place['rows'].append({**{name: clean(value) for name, value in row.items()},
                              'Localidade/Igreja': locality, '_written': written, '_line': number})

    churches = []
    for (country, _), place in places.items():
        rows = place['rows']
        code = COUNTRY_CODES.get(country)
        if not code:
            raise SystemExit(f'País sem código: {country!r}. Acrescente-o a COUNTRY_CODES.')
        locality = None if place['online'] else majority([row['Localidade/Igreja'] for row in rows])
        record_id = f'online-{slug(country)}' if place['online'] else f'{code.lower()}-{slug(locality)}'

        region_raw = majority([row['Província/Estado'] for row in rows])
        region = REGION_FIXES.get(region_raw, region_raw)
        if region != region_raw:
            notes['regions'].append(f'{record_id}: região «{region_raw}» escrita «{region}».')

        # An address is the workbook's own text, and only when it says more
        # than the place's name — as the row wrote it, or as it was joined.
        # Several different ones are reported, and the longest is kept.
        def says_more(row):
            text = fold(row['Endereço/Local do culto'])
            return text and text not in (fold(row['_written']), fold(row['Localidade/Igreja']))

        addresses = sorted({row['Endereço/Local do culto'] for row in rows if not place['online'] and says_more(row)},
                           key=len, reverse=True)
        if len(addresses) > 1:
            notes['addresses'].append(f'{record_id}: {len(addresses)} endereços diferentes; ficou o mais completo. {addresses}')
        address = addresses[0] if addresses else None

        # Responsibles, highest rank first. One person is one name or one
        # number: "Marcelo Cassange" and "Marcelo Cassengue" share a phone,
        # and Wilson Matuca is listed with two phones under one name.
        people = []
        for row in rows:
            name = row['Responsável']
            if not name:
                continue
            key = re.sub(r' (dos|das|da|de|do) ', ' ', f' {fold(name)} ').strip()
            phone = digits(row['Contacto'])
            person = next((item for item in people if key in item['keys'] or (phone and phone in item['digits'])), None)
            if person is None:
                person = {'keys': set(), 'digits': set(), 'names': [], 'ranks': [], 'phones': []}
                people.append(person)
            person['keys'].add(key)
            if phone:
                person['digits'].add(phone)
            person['names'].append(name)
            person['ranks'].append(rank_of(row['Cargo']))
            person['phones'].append((phone, row['Contacto'], fold(row['Dia'])))
        leaders = []
        for person in people:
            name = majority(person['names'])
            if len(set(person['names'])) > 1:
                notes['names'].append(f'{record_id}: {sorted(set(person["names"]))} — ficou «{name}».')
            by_digits = [phone[0] for phone in person['phones'] if phone[0]]
            sunday = next((phone[0] for phone in person['phones'] if phone[2] == 'domingo'), None)
            chosen = majority(by_digits, prefer=sunday)
            if len(set(by_digits)) > 1:
                notes['phones'].append(f'{record_id}: {name} com {sorted(set(phone[1] for phone in person["phones"]))} — ficou o que aparece em mais linhas.')
            sunday_written = next((phone[1] for phone in person['phones'] if phone[2] == 'domingo' and phone[0] == chosen), None)
            written = majority([phone[1] for phone in person['phones'] if phone[0] == chosen], prefer=sunday_written)
            leaders.append({'rank': majority(person['ranks']), 'name': name, 'phone': written})
        leaders.sort(key=lambda leader: RANK_ORDER.index(leader['rank']))
        if len(leaders) > 1:
            names = ', '.join(leader['rank'] + ' ' + leader['name'] for leader in leaders)
            notes['leaders'].append(f'{record_id}: {len(leaders)} responsáveis — {names}.')

        services = {}
        for row in rows:
            weekday = WEEKDAYS.get(fold(row['Dia']))
            if weekday is None:
                continue
            start = row['Horário']
            label = 'Via Zoom' if fold(row['Observações']) == 'via zoom' else None
            services[(weekday, start)] = {'weekday': weekday, 'start_time': start, 'label': label}

        churches.append({
            'record_id': record_id,
            'former_record_ids': [],
            'modality': 'online' if place['online'] else 'physical',
            'seat': SEATS.get(record_id),
            'country_code': code,
            'country': country,
            'region': region,
            'locality': locality,
            'address': address,
            'leader_name': f'{leaders[0]["rank"]} {leaders[0]["name"]}' if leaders else None,
            'leader_phone': leaders[0]['phone'] if leaders else None,
            'other_leaders': [{'name': f'{leader["rank"]} {leader["name"]}', 'phone': leader['phone']} for leader in leaders[1:]],
            'source': SOURCE,
            'verification_status': 'needs_review',
            'services': sorted(services.values(), key=lambda service: ((service['weekday'] + 6) % 7, service['start_time'] or '')),
            'lines': [row['_line'] for row in rows],
        })

    ids = Counter(church['record_id'] for church in churches)
    repeated = [record_id for record_id, count in ids.items() if count > 1]
    if repeated:
        raise SystemExit(f'Identificadores repetidos: {repeated}')
    missing = set(SEATS) - set(ids)
    if missing:
        raise SystemExit(f'SEATS aponta para lugares que não existem: {missing}')
    return churches, notes


def attach_former_ids(churches):
    """The ids the app used before this import, so old links and "A minha
    ISTN" keep finding the place (see former_record_ids in the migration)."""
    physical = json.loads((DATA / 'church-service-source-records.json').read_text('utf-8'))
    online = json.loads((DATA / 'online-communities-source-records.json').read_text('utf-8'))
    by_key = {}
    for church in churches:
        key = (church['country_code'], 'online') if church['modality'] == 'online' else (church['country_code'], fold(church['locality']))
        by_key[key] = church
    unmatched = []
    for record in physical:
        country = next((name for name, code in COUNTRY_CODES.items() if code == record['country_code']), None)
        locality = SAME_PLACE.get((country, record['locality']), record['locality'])
        church = by_key.get((record['country_code'], fold(locality)))
        if church:
            church['former_record_ids'].append(record['record_id'])
        else:
            unmatched.append(f'{record["record_id"]} · {record["locality"]} ({record["country_code"]})')
    for index, record in enumerate(online, start=1):
        church = by_key.get((COUNTRY_CODES.get(record['country']), 'online'))
        record_id = f'online_record_{index:03d}'
        if church:
            church['former_record_ids'].append(record_id)
        else:
            unmatched.append(f'{record_id} · {record["country"]}')
    return unmatched


# ------------------------------------------------------------------ saída --

def migration_sql(churches, name):
    """The migration: the places as one JSON document inside the import block.

    Not a temporary table, and not one INSERT per place: the Supabase SQL
    editor runs each statement of a file on its own, so anything the import
    needs has to travel inside the single statement that does it.
    """
    places = [{key: church[key] for key in (
        'record_id', 'former_record_ids', 'modality', 'seat', 'country_code', 'country', 'region', 'locality',
        'address', 'leader_name', 'leader_phone', 'other_leaders', 'source')}
        | {'services': [{key: service[key] for key in ('weekday', 'start_time', 'label')} for service in church['services']]}
        for church in churches]
    document = '[\n' + ',\n'.join('  ' + json.dumps(place, ensure_ascii=False) for place in places) + '\n]'
    if '$places$' in document:
        raise SystemExit('Os dados contêm $places$, a marca que delimita o documento JSON.')
    template = (ROOT / 'scripts' / 'import_churches.sql').read_text('utf-8')
    return template.replace('{{NAME}}', name).replace('{{PLACES}}', document).replace('{{COUNT}}', str(len(churches)))


def report(churches, notes, unmatched, workbook):
    physical = [church for church in churches if church['modality'] == 'physical']
    online = [church for church in churches if church['modality'] == 'online']
    lines = [
        '# Importação da lista geral de cultos',
        '',
        f'Gerado por `scripts/import_churches.py` a partir de `{Path(workbook).name}` em {date.today().isoformat()}.',
        '',
        f'{sum(len(church["lines"]) for church in churches)} linhas da folha → **{len(churches)} lugares**: '
        f'{len(physical)} presenciais e {len(online)} igrejas online, em {len({church["country"] for church in churches})} países.',
        '',
        'Os endereços são o texto da própria folha. Um lugar cujo «endereço» é apenas o seu nome fica sem endereço, em vez de lhe ser atribuído um.',
        '',
    ]

    def section(title, intro, items):
        lines.extend([f'## {title}', '', intro, ''])
        lines.extend([f'- {item}' for item in items] or ['_Nada a assinalar._'])
        lines.append('')

    section('Sedes — proposta a confirmar',
            'A folha não diz qual é a sede de cada país. Proposta: a sede mundial em Kifica; em cada país, o lugar do responsável de função mais alta (o primeiro da lista, em caso de empate). Muda-se no painel, em «Diretório».',
            [f'**{"Sede mundial" if church["seat"] == "mundial" else "Sede nacional"}** — {church["country"]}: {church["locality"]} · {church["leader_name"]}'
             for church in churches if church['seat']])
    section('Lugares juntados', 'Linhas que a folha escreve com dois nomes e que são o mesmo lugar: o mesmo responsável, o mesmo número.', notes['merged'])
    section('Mais do que um responsável', 'Aparecem todos na página da igreja; o primeiro é o contacto principal.', notes['leaders'])
    section('Números diferentes para a mesma pessoa', 'Ficou o número que aparece em mais linhas (o de domingo, em caso de empate). Confirme com o responsável.', notes['phones'])
    section('Grafias diferentes do mesmo nome', 'Ficou a grafia mais frequente.', notes['names'])
    section('Regiões corrigidas', 'Erros de digitação e uma abreviatura.', notes['regions'])
    section('Endereços diferentes para o mesmo lugar', 'Ficou o mais completo.', notes['addresses'])
    section('Com endereço', 'Lugares cuja folha indica um endereço.',
            [f'{church["locality"]} ({church["country"]}): {church["address"]}' for church in physical if church['address']])
    section('Sem endereço na folha', 'Aparecem com «Endereço a confirmar com o responsável».',
            [f'{church["locality"]} ({church["region"] or church["country"]})' for church in physical if not church['address']])
    section('Registos anteriores sem correspondência', 'Estavam na base de dados e não aparecem na folha. A migração não os apaga: confirme se ainda existem.', unmatched)
    return '\n'.join(lines).rstrip() + '\n'


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('workbook')
    parser.add_argument('--migracao', help='caminho do ficheiro SQL a escrever em supabase/migrations/')
    args = parser.parse_args()

    churches, notes = build(args.workbook)
    unmatched = attach_former_ids(churches)
    public = [{key: value for key, value in church.items() if key != 'lines'} for church in churches]

    (DATA / 'igrejas.json').write_text(json.dumps({
        'source': SOURCE,
        'workbook': Path(args.workbook).name,
        'importedOn': date.today().isoformat(),
        'churches': public,
    }, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (DATA / 'importacao-igrejas.md').write_text(report(churches, notes, unmatched, args.workbook), encoding='utf-8')
    print(f'{len(churches)} lugares -> data/igrejas.json e data/importacao-igrejas.md')

    if args.migracao:
        target = Path(args.migracao)
        if not target.is_absolute():
            target = ROOT / target
        target.write_text(migration_sql(public, target.stem), encoding='utf-8')
        print(f'Migração -> {target.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
